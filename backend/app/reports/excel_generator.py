from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timezone
from typing import Dict, List, Any

def generate_excel(report_title: str, records: list[dict]) -> bytes:
    """Generate a single-sheet Excel (xlsx) report for backwards compatibility."""
    wb = Workbook()
    ws = wb.active
    ws.title = report_title[:31]
    
    # Header row styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    if records:
        headers = list(records[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for rec in records:
            ws.append([rec.get(h, "") for h in headers])

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
    else:
        ws.append(["No data available."])

    ws.append([])
    ws.append([f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')} UTC"])

    buffer = BytesIO()
    wb.save(buffer)
    excel_data = buffer.getvalue()
    buffer.close()
    return excel_data


def generate_multi_sheet_excel(sheets_data: Dict[str, List[Dict[str, Any]]]) -> bytes:
    """
    Generates a professional multi-sheet Excel workbook.
    sheets_data: Dict mapping sheet_name -> list of row dicts.
    Only included sheets that have rows.
    """
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for sheet_name, rows in sheets_data.items():
        if not rows:
            continue
        ws = wb.create_sheet(title=sheet_name[:31])
        headers = list(rows[0].keys())
        ws.append(headers)

        # Style header
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r in rows:
            ws.append([r.get(h, "") for h in headers])

        # Column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 40)

        # Add timestamp note
        ws.append([])
        ws.append([f"Report generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')} UTC"])

    if not wb.sheetnames:
        ws = wb.create_sheet(title="Executive Summary")
        ws.append(["No department datasets have been uploaded yet."])

    buffer = BytesIO()
    wb.save(buffer)
    excel_data = buffer.getvalue()
    buffer.close()
    return excel_data
